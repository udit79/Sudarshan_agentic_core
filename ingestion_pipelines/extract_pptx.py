"""PPTX text and notes extractor using python-pptx.

Follows the "SLIDES + NOTES" design:
- Iterates slides in sequence
- Extracts text from all text frames and table cells
- Extracts speaker notes if present
- Preserves structure with clear slide delimiters
"""

from __future__ import annotations

import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
from pptx import Presentation


def _ppt_master_root() -> Path | None:
    configured = os.getenv("PPT_MASTER_ROOT", "").strip()
    if not configured or os.getenv("PPT_MASTER_INTAKE_ENABLED", "false").lower() not in {"1", "true", "yes"}:
        return None
    root = Path(configured).expanduser()
    return root if root.is_dir() else None


def _extract_with_ppt_master(file_path: str, root: Path) -> str | None:
    """Run PPT Master's optional source converters as an enrichment boundary."""

    skill_dir = root / "skills" / "ppt-master"
    converter = skill_dir / "scripts" / "source_to_md.py"
    intake = skill_dir / "scripts" / "pptx_intake.py"
    if not converter.is_file():
        return None
    with tempfile.TemporaryDirectory(prefix="sudarshan-ppt-intake-") as temp_dir:
        output_dir = Path(temp_dir)
        completed = subprocess.run(
            [sys.executable, str(converter), str(file_path), "-o", str(output_dir)],
            capture_output=True,
            text=True,
            timeout=180,
        )
        if completed.returncode != 0:
            raise RuntimeError(completed.stderr.strip() or "PPT Master source conversion failed")
        markdown_files = sorted(output_dir.rglob("*.md"))
        if not markdown_files:
            raise RuntimeError("PPT Master source conversion returned no Markdown file")
        sections = [markdown_files[0].read_text(encoding="utf-8", errors="replace")]
        if intake.is_file():
            analysis_dir = output_dir / "analysis"
            analysis = subprocess.run(
                [sys.executable, str(intake), str(file_path), "-o", str(analysis_dir)],
                capture_output=True,
                text=True,
                timeout=180,
            )
            if analysis.returncode != 0:
                raise RuntimeError(analysis.stderr.strip() or "PPT Master PPTX intake failed")
            profile = analysis_dir / "source_profile.json"
            if profile.is_file():
                try:
                    sections.append("[PPT Master source profile]\n" + json.dumps(
                        json.loads(profile.read_text(encoding="utf-8")),
                        ensure_ascii=False,
                        indent=2,
                    ))
                except (json.JSONDecodeError, UnicodeDecodeError) as exc:
                    raise RuntimeError("PPT Master source profile was invalid JSON") from exc
        return "\n\n".join(sections).strip()


def extract_text_from_pptx(file_path: str) -> str:
    """Extracts slide content and speaker notes from a PowerPoint presentation."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Source file not found: {file_path}")

    if os.getenv("PPT_MASTER_INTAKE_ENABLED", "false").lower() in {"1", "true", "yes"}:
        try:
            from integrations.providers.pptmaster.intake import PPTMasterIntake
            intake = PPTMasterIntake()
            profile = intake.extract(str(path))
            return profile.markdown_content
        except Exception:
            if os.getenv("PPT_MASTER_INTAKE_STRICT", "false").lower() in {"1", "true", "yes"}:
                raise

    try:
        prs = Presentation(str(path))
    except Exception as err:
        raise ValueError(f"Corrupted or unreadable PowerPoint presentation '{path.name}': {err}") from err
    slide_chunks: list[str] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_lines: list[str] = []

        # 1. Collect text from all shapes (text frames and tables)
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_lines.append(text)
            elif shape.has_table:
                table_lines: list[str] = []
                for row in shape.table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        table_lines.append(" | ".join(row_cells))
                if table_lines:
                    slide_lines.extend(table_lines)

        # 2. Collect speaker notes if present
        notes_text = ""
        if slide.has_notes_slide:
            notes_slide = slide.notes_slide
            if notes_slide.notes_text_frame:
                notes_text = notes_slide.notes_text_frame.text.strip()

        # 3. Assemble slide representation
        content_block = "\n".join(slide_lines).strip()
        header = f"--- Slide {slide_num} ---"

        parts: list[str] = [header]
        if content_block:
            parts.append(content_block)
        if notes_text:
            parts.append(f"[Speaker Notes]:\n{notes_text}")

        # If the slide had either content or notes, record it
        if content_block or notes_text:
            slide_chunks.append("\n".join(parts))
        else:
            slide_chunks.append(f"{header}\n(Empty Slide)")

    return "\n\n".join(slide_chunks).strip()
