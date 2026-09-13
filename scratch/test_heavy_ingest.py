import os
import sys
import tempfile
import time
from pathlib import Path

# Add core to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches, Pt
from ingestion_pipelines import ingest_file
from ingestion_pipelines.extract_pdf import extract_text_from_pdf
from ingestion_pipelines.extract_pptx import extract_text_from_pptx

def test_large_pdf_handling():
    print("=== Testing Heavy PDF Generation & Extraction ===")
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_pdf = f.name

    try:
        # Create a 25-page digital PDF
        doc = fitz.open()
        for i in range(1, 26):
            page = doc.new_page(width=595, height=842)
            page.insert_text(
                (50, 72),
                f"SUDARSHAN TACTICAL INTEL REPORT - PAGE {i}\n"
                f"Classification: RESTRICTED // NTRO REL TO INDIA\n"
                f"Timestamp: 2026-09-13T16:00:00Z\n"
                f"Section {i}.1: Overview of border surveillance sector {i}.\n"
                f"Payload telemetry: latitude=34.{i}000, longitude=74.{i}000.\n"
                f"Detailed observation notes: Sector {i} reports nominal operation with perimeter drone patrol active."
            )
        doc.save(temp_pdf)
        doc.close()
        file_size_kb = os.path.getsize(temp_pdf) / 1024
        print(f"[PDF] Created 25-page PDF ({file_size_kb:.2f} KB)")

        start_time = time.perf_counter()
        extracted = extract_text_from_pdf(temp_pdf)
        duration = time.perf_counter() - start_time
        print(f"[PDF] Extraction completed in {duration:.3f}s. Extracted {len(extracted)} chars across 25 pages.")
        assert "--- Page 1 ---" in extracted
        assert "--- Page 25 ---" in extracted

        # Test Windows unlinking immediately to verify no file handle leak
        Path(temp_pdf).unlink()
        print("[PDF] File handle test: Successfully deleted temp PDF without Windows lock error!")
    except Exception as e:
        print(f"[PDF ERROR] {e}")
        Path(temp_pdf).unlink(missing_ok=True)
        raise

def test_large_pptx_handling():
    print("\n=== Testing Heavy PPTX Generation & Extraction ===")
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as f:
        temp_pptx = f.name

    try:
        prs = Presentation()
        # Create 20 slides with text, tables, and notes
        for i in range(1, 21):
            slide_layout = prs.slide_layouts[6]  # Blank
            slide = prs.slides.add_slide(slide_layout)
            
            # Text box
            tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
            tf = tx_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"Slide {i}: Tactical Briefing Sector {i}"
            
            # Add table
            table_shape = slide.shapes.add_table(3, 3, Inches(1), Inches(2.5), Inches(6), Inches(2))
            table = table_shape.table
            table.cell(0, 0).text = "Unit ID"
            table.cell(0, 1).text = "Status"
            table.cell(0, 2).text = "Coordinates"
            table.cell(1, 0).text = f"ALPHA-{i}"
            table.cell(1, 1).text = "PATROLLING"
            table.cell(1, 2).text = f"GRID-34-{i}"

            # Add speaker notes
            notes_slide = slide.notes_slide
            tf_notes = notes_slide.notes_text_frame
            tf_notes.text = f"Speaker guidance for slide {i}: Emphasize speed and reliability under network degradation."

        prs.save(temp_pptx)
        file_size_kb = os.path.getsize(temp_pptx) / 1024
        print(f"[PPTX] Created 20-slide presentation with tables & notes ({file_size_kb:.2f} KB)")

        start_time = time.perf_counter()
        extracted = extract_text_from_pptx(temp_pptx)
        duration = time.perf_counter() - start_time
        print(f"[PPTX] Extraction completed in {duration:.3f}s. Extracted {len(extracted)} chars across 20 slides.")
        assert "Slide 1:" in extracted
        assert "ALPHA-1" in extracted
        assert "[Speaker Notes]:" in extracted
        assert "Speaker guidance for slide 20" in extracted

        Path(temp_pptx).unlink()
        print("[PPTX] File handle test: Successfully deleted temp PPTX without Windows lock error!")
    except Exception as e:
        print(f"[PPTX ERROR] {e}")
        Path(temp_pptx).unlink(missing_ok=True)
        raise

def test_end_to_end_heavy_pipeline():
    print("\n=== Testing End-to-End Heavy Ingestion Pipeline (Blocks & Indexing) ===")
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        temp_pdf = f.name

    try:
        doc = fitz.open()
        for i in range(1, 31):
            page = doc.new_page()
            page.insert_text((50, 72), f"Paragraph {i}: Detailed tactical observation on radar grid {i}.\n" * 5)
        doc.save(temp_pdf)
        doc.close()

        start_time = time.perf_counter()
        ingested = ingest_file(temp_pdf, user_id="heavy_tester", case_id="case_heavy_01", task_id="task_heavy_01")
        duration = time.perf_counter() - start_time
        print(f"[E2E Ingest] Ingested 30-page PDF into {len(ingested.evidence_blocks)} evidence blocks in {duration:.3f}s")
        assert len(ingested.evidence_blocks) == 30
        print(f"[E2E Ingest] Verified all 30 page blocks generated properly.")
        Path(temp_pdf).unlink()
    except Exception as e:
        Path(temp_pdf).unlink(missing_ok=True)
        raise

if __name__ == "__main__":
    test_large_pdf_handling()
    test_large_pptx_handling()
    test_end_to_end_heavy_pipeline()
    print("\n[ALL DIAGNOSTIC BENCHMARKS PASSED]")

