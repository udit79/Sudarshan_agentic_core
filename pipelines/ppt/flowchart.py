"""Deterministic flowchart geometry and dual SVG/editable-PPTX rendering."""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from pipelines.ppt.schemas import FlowchartEdge, FlowchartNode, FlowchartSpec, LayoutBox, VisualIR


@dataclass(frozen=True, slots=True)
class LayoutNode:
    node: FlowchartNode
    box: LayoutBox


@dataclass(frozen=True, slots=True)
class FlowchartLayout:
    spec: FlowchartSpec
    nodes: tuple[LayoutNode, ...]

    @property
    def by_id(self) -> dict[str, LayoutNode]:
        return {item.node.node_id: item for item in self.nodes}


@dataclass(frozen=True, slots=True)
class FlowchartRenderArtifact:
    path: str
    svg: str
    node_count: int
    edge_count: int


def flowchart_from_visual_ir(visual: VisualIR) -> FlowchartSpec:
    """Convert a typed visual IR payload into the strict flowchart graph IR."""

    if visual.kind != "flowchart":
        raise ValueError("visual IR kind must be flowchart")
    data = dict(visual.data)
    raw_nodes = data.get("nodes", [])
    raw_edges = data.get("edges", [])
    if not isinstance(raw_nodes, list) or not isinstance(raw_edges, list):
        raise ValueError("flowchart visual data requires nodes and edges lists")
    nodes = []
    for raw in raw_nodes:
        if not isinstance(raw, dict):
            raise ValueError("flowchart nodes must be objects")
        nodes.append(
            FlowchartNode(
                node_id=str(raw.get("node_id", raw.get("id", ""))),
                label=str(raw.get("label", "")),
                kind=str(raw.get("kind", "process")),
                group=raw.get("group"),
            )
        )
    edges = []
    for raw in raw_edges:
        if not isinstance(raw, dict):
            raise ValueError("flowchart edges must be objects")
        edges.append(
            FlowchartEdge(
                source=str(raw.get("source", raw.get("from", ""))),
                target=str(raw.get("target", raw.get("to", ""))),
                label=str(raw.get("label", "")),
            )
        )
    return FlowchartSpec(
        flowchart_id=visual.visual_id,
        direction=str(data.get("direction", "left-to-right")),
        nodes=nodes,
        edges=edges,
        layout_hints=dict(data.get("layout_hints", {})),
        style_tokens=dict(data.get("style_tokens", {})),
    )


def validate_flowchart(spec: FlowchartSpec) -> None:
    """Validate DAG structure and reject isolated nodes before rendering."""

    incoming = {node.node_id: 0 for node in spec.nodes}
    outgoing = {node.node_id: [] for node in spec.nodes}
    for edge in spec.edges:
        incoming[edge.target] += 1
        outgoing[edge.source].append(edge.target)

    if len(spec.nodes) > 1 and any(not outgoing[node_id] and incoming[node_id] == 0 for node_id in incoming):
        raise ValueError("flowchart contains isolated nodes")

    queue = [node_id for node_id, count in incoming.items() if count == 0]
    visited = 0
    indegree = dict(incoming)
    while queue:
        current = queue.pop(0)
        visited += 1
        for target in outgoing[current]:
            indegree[target] -= 1
            if indegree[target] == 0:
                queue.append(target)
    if visited != len(spec.nodes):
        raise ValueError("flowchart must be acyclic")


def layout_flowchart(spec: FlowchartSpec, bounds: LayoutBox | None = None) -> FlowchartLayout:
    """Lay out a bounded DAG in deterministic layers using normalized coords."""

    validate_flowchart(spec)
    bounds = bounds or LayoutBox(x=0.08, y=0.18, width=0.84, height=0.64)
    incoming = {node.node_id: [] for node in spec.nodes}
    outgoing = {node.node_id: [] for node in spec.nodes}
    for edge in spec.edges:
        outgoing[edge.source].append(edge.target)
        incoming[edge.target].append(edge.source)

    levels: dict[str, int] = {}
    pending = [node.node_id for node in spec.nodes if not incoming[node.node_id]]
    while pending:
        current = pending.pop(0)
        level = levels.get(current, 0)
        for target in outgoing[current]:
            levels[target] = max(levels.get(target, 0), level + 1)
            if all(parent in levels for parent in incoming[target]):
                pending.append(target)

    grouped: dict[int, list[FlowchartNode]] = {}
    for node in spec.nodes:
        grouped.setdefault(levels.get(node.node_id, 0), []).append(node)
    level_count = max(grouped) + 1
    max_rows = max(len(nodes) for nodes in grouped.values())
    gap_x = min(0.04, bounds.width / max(1, level_count * 8))
    gap_y = min(0.04, bounds.height / max(1, max_rows * 8))
    node_width = min(0.22, (bounds.width - gap_x * max(0, level_count - 1)) / level_count * 0.78)
    node_height = min(0.14, (bounds.height - gap_y * max(0, max_rows - 1)) / max_rows * 0.72)

    result: list[LayoutNode] = []
    horizontal = spec.direction == "left-to-right"
    for level, nodes in sorted(grouped.items()):
        for row, node in enumerate(sorted(nodes, key=lambda item: item.node_id)):
            if horizontal:
                x = bounds.x + level * (bounds.width / level_count)
                y = bounds.y + (bounds.height - (len(nodes) * node_height + (len(nodes) - 1) * gap_y)) / 2
                y += row * (node_height + gap_y)
            else:
                y = bounds.y + level * (bounds.height / level_count)
                x = bounds.x + (bounds.width - (len(nodes) * node_width + (len(nodes) - 1) * gap_x)) / 2
                x += row * (node_width + gap_x)
            result.append(LayoutNode(node=node, box=LayoutBox(x=x, y=y, width=node_width, height=node_height)))
    return FlowchartLayout(spec=spec, nodes=tuple(result))


def render_flowchart_svg(
    spec: FlowchartSpec,
    *,
    width: int = 1200,
    height: int = 675,
    layout: FlowchartLayout | None = None,
) -> str:
    """Render the same layout to an inspectable SVG preview."""

    layout = layout or layout_flowchart(spec)
    by_id = layout.by_id
    colors = {"start": "#D1FAE5", "process": "#DBEAFE", "decision": "#FEF3C7", "review": "#EDE9FE", "end": "#FEE2E2"}
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" role="img" aria-labelledby="diagram-title diagram-desc" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        f'<title id="diagram-title">{escape(spec.flowchart_id)}</title>',
        f'<desc id="diagram-desc">Static flowchart {escape(spec.flowchart_id)} with {len(spec.nodes)} nodes and {len(spec.edges)} connections.</desc>',
        '<defs><marker id="arrow" markerWidth="10" markerHeight="10" refX="8" refY="3" orient="auto"><path d="M0,0 L0,6 L9,3 z" fill="#475569"/></marker></defs>',
        '<rect width="100%" height="100%" fill="#F8FAFC"/>',
    ]
    for edge in spec.edges:
        source = by_id[edge.source].box
        target = by_id[edge.target].box
        x1, y1, x2, y2 = _edge_points(source, target, width, height)
        parts.append(f'<path d="M{x1:.1f},{y1:.1f} L{x2:.1f},{y2:.1f}" fill="none" stroke="#64748B" stroke-width="2" marker-end="url(#arrow)"/>')
        if edge.label:
            parts.append(f'<text x="{(x1+x2)/2:.1f}" y="{(y1+y2)/2-5:.1f}" text-anchor="middle" font-family="Arial,sans-serif" font-size="14" fill="#334155">{escape(edge.label)}</text>')
    for item in layout.nodes:
        box = item.box
        x, y, w, h = box.x * width, box.y * height, box.width * width, box.height * height
        fill = colors.get(item.node.kind, "#DBEAFE")
        parts.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{w:.1f}" height="{h:.1f}" rx="14" fill="{fill}" stroke="#1E3A8A" stroke-width="2"/>')
        parts.append(f'<text x="{x+w/2:.1f}" y="{y+h/2+5:.1f}" text-anchor="middle" font-family="Arial,sans-serif" font-size="18" font-weight="600" fill="#0F172A">{escape(item.node.label)}</text>')
    parts.append("</svg>")
    return "".join(parts)


def render_flowchart_pptx(
    spec: FlowchartSpec,
    output_path: str | Path,
    *,
    layout: FlowchartLayout | None = None,
    title: str | None = None,
) -> FlowchartRenderArtifact:
    """Create an editable PPTX slide from the shared flowchart layout."""

    layout = layout or layout_flowchart(spec)
    svg = render_flowchart_svg(spec, layout=layout)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    if title:
        textbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.2), Inches(12.1), Inches(0.45))
        textbox.text_frame.text = title
        textbox.text_frame.paragraphs[0].font.size = Pt(24)
        textbox.text_frame.paragraphs[0].font.bold = True

    by_id = layout.by_id
    for edge in spec.edges:
        source = by_id[edge.source].box
        target = by_id[edge.target].box
        x1, y1, x2, y2 = _edge_points(source, target, 13.333, 7.5)
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        connector.line.color.rgb = RGBColor(100, 116, 139)
        connector.line.width = Pt(1.5)
        if edge.label:
            label = slide.shapes.add_textbox(Inches((x1+x2)/2 - 0.6), Inches((y1+y2)/2 - 0.15), Inches(1.2), Inches(0.3))
            label.text_frame.text = edge.label
            label.text_frame.paragraphs[0].font.size = Pt(9)
            label.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    fills = {"start": RGBColor(209, 250, 229), "process": RGBColor(219, 234, 254), "decision": RGBColor(254, 243, 199), "review": RGBColor(237, 233, 254), "end": RGBColor(254, 226, 226)}
    for item in layout.nodes:
        box = item.box
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(box.x * 13.333),
            Inches(box.y * 7.5),
            Inches(box.width * 13.333),
            Inches(box.height * 7.5),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fills.get(item.node.kind, fills["process"])
        shape.line.color.rgb = RGBColor(30, 58, 138)
        shape.text_frame.clear()
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.text = item.node.label
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.font.size = Pt(16)
        paragraph.font.bold = True

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output))
    return FlowchartRenderArtifact(path=str(output), svg=svg, node_count=len(layout.nodes), edge_count=len(spec.edges))


def _edge_points(source: LayoutBox, target: LayoutBox, width: float, height: float) -> tuple[float, float, float, float]:
    """Connect nearest box faces in output units."""

    sx, sy = (source.x + source.width / 2) * width, (source.y + source.height / 2) * height
    tx, ty = (target.x + target.width / 2) * width, (target.y + target.height / 2) * height
    if abs(tx - sx) >= abs(ty - sy):
        return (
            (source.x + source.width) * width if tx >= sx else source.x * width,
            sy,
            target.x * width if tx >= sx else (target.x + target.width) * width,
            ty,
        )
    return (
        sx,
        (source.y + source.height) * height if ty >= sy else source.y * height,
        tx,
        target.y * height if ty >= sy else (target.y + target.height) * height,
    )


__all__ = [
    "FlowchartLayout",
    "FlowchartRenderArtifact",
    "flowchart_from_visual_ir",
    "LayoutNode",
    "layout_flowchart",
    "render_flowchart_pptx",
    "render_flowchart_svg",
    "validate_flowchart",
]
