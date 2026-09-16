"""PPTX artifact renderer for the PPT generation pipeline.

Converts a validated PresentationOutput into a .pptx file using python-pptx.
Written to artifacts/presentations/ with a stable, auditable filename.
"""

from __future__ import annotations

import shutil
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, NamedTuple
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from pipelines.ppt.schemas import (
    DeckManifest,
    IncrementalDeckEdit,
    PresentationOutput,
    PresentationTheme,
    SlideContent,
    resolve_presentation_theme,
)
from pipelines.ppt.template_workspace import PptTemplateContract


_ARTIFACT_DIR = Path("artifacts") / "presentations"


class PresentationArtifact(NamedTuple):
    path: str
    presentation_id: str
    slide_count: int
    content: str   # human-readable summary for memory write-back
    deck_manifest: dict[str, Any] | None = None
    theme: dict[str, Any] | None = None


def _resolve_base_artifact(artifact_id: str) -> Path:
    """Mock resolution of a control-plane artifact ID into a local path."""
    path = Path(artifact_id)
    if not path.is_file():
        raise FileNotFoundError(f"Base artifact not found locally: {artifact_id}")
    return path


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.removeprefix("#"))


def _style_text_frame(text_frame: Any, *, theme: PresentationTheme, size: int, bold: bool = False) -> None:
    for paragraph in text_frame.paragraphs:
        paragraph.font.name = theme.body_font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = _rgb(theme.foreground)


def _render_slide(prs: Presentation, slide_data: SlideContent, theme: PresentationTheme) -> None:
    """Render a slide based on its explicit layout type."""
    # Simplified layout mapping
    if slide_data.layout == "cover":
        layout = prs.slide_layouts[0]
    elif slide_data.layout in ("agenda", "conclusion", "closing", "content", "two_column"):
        layout = prs.slide_layouts[1]
    else:
        layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(theme.background)

    # A small native accent bar makes the selected palette observable in the
    # artifact and gives the QA gate a stable, editable layer to inspect.
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
    accent.name = "theme.accent"
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(theme.accent)
    accent.line.fill.background()

    # Title
    if slide.shapes.title:
        slide.shapes.title.text = slide_data.title
        title_frame = slide.shapes.title.text_frame
        _style_text_frame(title_frame, theme=theme, size=36 if slide_data.layout == "cover" else 24, bold=True)
        for paragraph in title_frame.paragraphs:
            paragraph.font.name = theme.title_font

    # Body (Bullets)
    if len(slide.placeholders) > 1:
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()
        for i, bullet in enumerate(slide_data.bullets):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = bullet
            para.font.name = theme.body_font
            para.font.size = Pt(18 if slide_data.layout in ("cover", "agenda", "conclusion") else 16)
            para.font.color.rgb = _rgb(theme.foreground)
            para.level = 0

    # Speaker notes
    if slide_data.speaker_notes:
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.text = slide_data.speaker_notes


def _update_slide_in_place(prs: Presentation, slide_index: int, slide_patch: dict) -> None:
    """Apply a semantic update to an existing slide."""
    slide = prs.slides[slide_index]
    content = slide_patch.get("content", {})
    
    if slide_patch.get("operation") == "replace_content":
        if "title" in content and slide.shapes.title:
            slide.shapes.title.text = content["title"]
        if "bullets" in content and len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.clear()
            for i, bullet in enumerate(content["bullets"]):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = bullet
    elif slide_patch.get("operation") == "rewrite_notes":
        if "speaker_notes" in content:
            notes_frame = slide.notes_slide.notes_text_frame
            notes_frame.text = content["speaker_notes"]


def _stable_hash(value: Any) -> str:
    payload = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _slide_source_hash(slide: SlideContent | dict[str, Any]) -> str:
    payload = slide.model_dump(mode="json") if hasattr(slide, "model_dump") else dict(slide)
    return _stable_hash(payload)


def _slide_part_hashes(path: str | Path) -> dict[int, str]:
    """Hash slide XML parts by their one-based PowerPoint slide number."""

    hashes: dict[int, str] = {}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("ppt/slides/slide") or not name.endswith(".xml"):
                continue
            number = name.removeprefix("ppt/slides/slide").removesuffix(".xml")
            if number.isdigit():
                hashes[int(number)] = hashlib.sha256(archive.read(name)).hexdigest()
    return hashes


def _descendant_closure(slide_ids: set[str], manifest: DeckManifest) -> set[str]:
    """Return all downstream slides affected by a changed slide."""

    closure = set(slide_ids)
    changed = True
    while changed:
        changed = False
        for slide in manifest.slides:
            if slide.slide_id not in closure and any(dep in closure for dep in slide.dependencies):
                closure.add(slide.slide_id)
                changed = True
    return closure


def render_presentation(
    output: PresentationOutput | None = None,
    *,
    incremental_edit: IncrementalDeckEdit | None = None,
    previous_manifest: DeckManifest | None = None,
    approved_by: str | None = None,
    theme: PresentationTheme | None = None,
    template_contract: PptTemplateContract | None = None,
) -> PresentationArtifact:
    """Render PresentationOutput → PPTX file → PresentationArtifact.
    
    If `incremental_edit` is provided, performs a targeted semantic update to the existing presentation.
    """

    if output is None and incremental_edit is None:
        raise ValueError("Must provide either output or incremental_edit")

    theme = theme or resolve_presentation_theme()

    if output is not None and template_contract is not None:
        if output.template_id != template_contract.template_id:
            raise ValueError("presentation template_id does not match template_contract")
        template_contract.validate_rendered_layouts([_template_purpose(slide.layout) for slide in output.slides])
    if output is not None and output.template_id != "native-default" and template_contract is None:
        raise ValueError("non-default presentation templates require a validated template_contract")
    
    if incremental_edit is not None:
        if previous_manifest is None:
            raise ValueError("Incremental edits require a previous_manifest")
        if previous_manifest.deck_id != incremental_edit.base_manifest_id:
            raise ValueError("Stale or mismatched base_manifest_id")
            
        presentation_id = previous_manifest.deck_id
    else:
        presentation_id = output.presentation_id

    _ARTIFACT_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    safe_id = presentation_id.replace("/", "-").replace(" ", "_")[:60]

    if incremental_edit:
        next_ver = previous_manifest.version + 1
        filename = f"{safe_id}_v{next_ver}_{timestamp}.pptx"
    else:
        filename = f"{safe_id}_{timestamp}.pptx"

    file_path = _ARTIFACT_DIR / filename

    if incremental_edit:
        # INCREMENTAL UPDATE
        base_path = _resolve_base_artifact(incremental_edit.base_artifact_id)
        
        # Check for unsupported structural edits
        if any(patch.operation not in ("replace_content", "replace_visual", "rewrite_notes") for patch in incremental_edit.edits):
            raise ValueError("Structural edits (additions/deletions) are not supported incrementally.")

        changed_ids = {patch.slide_id for patch in incremental_edit.edits}
        patch_source_hashes = {
            patch.slide_id: _stable_hash(patch.model_dump(mode="json"))
            for patch in incremental_edit.edits
        }
        requested_scope = set(incremental_edit.requested_scope) or set(changed_ids)
        if not changed_ids.issubset(requested_scope):
            raise ValueError("incremental edits contain slides outside requested_scope")
        affected_ids = _descendant_closure(requested_scope, previous_manifest)
        invalidated_ids = affected_ids - changed_ids
        before_hashes = _slide_part_hashes(base_path)
        
        # We copy the base path to avoid mutating it before a successful save
        shutil.copy2(base_path, file_path)
        prs = Presentation(file_path)

        # Apply patches
        for patch in incremental_edit.edits:
            if patch.slide_id in changed_ids:
                try:
                    # Resolve slide_id to index via manifest
                    slide_index = previous_manifest.slide_order.index(patch.slide_id)
                    _update_slide_in_place(prs, slide_index, patch.model_dump())
                except ValueError:
                    pass

        prs.save(str(file_path))
        after_hashes = _slide_part_hashes(file_path)

        for index, slide_id in enumerate(previous_manifest.slide_order, start=1):
            if slide_id not in affected_ids and before_hashes.get(index) != after_hashes.get(index):
                raise ValueError(f"untouched slide XML changed during incremental update: {slide_id}")
        
        manifest_slides = []
        for s in previous_manifest.slides:
            if s.slide_id in changed_ids:
                status = "rendered"
            elif s.slide_id in invalidated_ids:
                status = "invalidated"
            else:
                status = "unchanged"
            index = previous_manifest.slide_order.index(s.slide_id) + 1
            manifest_slides.append({
                "slide_id": s.slide_id,
                "order": s.order,
                "status": status,
                "dependencies": s.dependencies,
                "source_ir_hash": patch_source_hashes.get(s.slide_id, s.source_ir_hash),
                "rendered_part_hash": after_hashes.get(index),
                "evidence_refs": s.evidence_refs,
                "asset_refs": s.asset_refs,
                "layout_id": s.layout_id,
            })
            
        deck_manifest = {
            "deck_id": presentation_id,
            "version": previous_manifest.version + 1,
            "base_artifact_id": str(file_path),
            "template_id": previous_manifest.template_id,
            "template_version": previous_manifest.template_version,
            "theme_id": previous_manifest.theme_id,
            "theme_hash": previous_manifest.theme_hash or theme.hash(),
            "theme_tokens": previous_manifest.theme_tokens or theme.model_dump(mode="json"),
            "master_id": previous_manifest.master_id,
            "slide_order": previous_manifest.slide_order,
            "slides": manifest_slides,
            "manifest_version": "1.0"
        }
        
        content_lines = [f"# Incremental Update to {presentation_id}"]
        slide_count = len(previous_manifest.slide_order)

    else:
        # FULL REBUILD
        prs = Presentation()
        # Standard widescreen (13.33 x 7.5 inches)
        prs.slide_width = Inches(13.33)
        prs.slide_height = Inches(7.5)

        for slide_data in output.slides:
            _render_slide(prs, slide_data, theme)

        prs.save(str(file_path))

        content_lines = [
            f"# {output.title}",
            f"Classification: {output.classification_level}",
            f"Distribution: {output.distribution}",
            "",
            "## Slides",
        ]
        for slide in output.slides:
            content_lines.append(f"### Slide {slide.order}: {slide.title}")
            content_lines.extend([f"- {b}" for b in slide.bullets])

        if approved_by:
            content_lines.append(f"\nApproved by: {approved_by}")

        manifest_slides = []
        rendered_hashes = _slide_part_hashes(file_path)
        for index, s in enumerate(output.slides, start=1):
            manifest_slides.append({
                "slide_id": s.slide_id,
                "order": s.order,
                "status": "rendered",
                "dependencies": [],
                "source_ir_hash": _slide_source_hash(s),
                "rendered_part_hash": rendered_hashes.get(index),
                "layout_id": s.layout,
            })
            
        deck_manifest = {
            "deck_id": presentation_id,
            "version": 1 if not previous_manifest else previous_manifest.version + 1,
            "base_artifact_id": str(file_path),
            "template_id": output.template_id,
            "template_version": output.template_version,
            "theme_id": "ntro-briefing",
            "theme_hash": theme.hash(),
            "theme_tokens": theme.model_dump(mode="json"),
            "master_id": template_contract.master_id if template_contract is not None else "native-master",
            "slide_order": [s["slide_id"] for s in manifest_slides],
            "slides": manifest_slides,
            "manifest_version": "1.0"
        }
        slide_count = len(output.slides)

    return PresentationArtifact(
        path=str(file_path),
        presentation_id=presentation_id,
        slide_count=slide_count,
        content="\n".join(content_lines),
        deck_manifest=deck_manifest,
        theme=theme.model_dump(mode="json"),
    )


def _template_purpose(layout: str) -> str:
    return {
        "cover": "title",
        "agenda": "narrative",
        "content": "narrative",
        "two_column": "narrative",
        "conclusion": "narrative",
        "closing": "closing",
    }.get(layout, "narrative")
