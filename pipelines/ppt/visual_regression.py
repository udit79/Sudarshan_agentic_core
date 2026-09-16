"""Small offline PPT visual-contract regression checks.

This does not pretend to replace PowerPoint rendering. It catches the cheap,
repeatable regressions first: slide count, native text layers, geometry, and
theme layer presence. Raster smoke tests can build on the same fixture later.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping
import hashlib
import json

from pptx import Presentation


def pptx_visual_snapshot(path: str | Path) -> dict[str, Any]:
    presentation = Presentation(str(path))
    slides: list[dict[str, Any]] = []
    for slide in presentation.slides:
        shapes: list[dict[str, Any]] = []
        for shape in slide.shapes:
            fill_color = None
            try:
                fill_color = str(shape.fill.fore_color.rgb)
            except (AttributeError, TypeError, ValueError):
                pass
            font_colors: list[str] = []
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        try:
                            if run.font.color.rgb is not None:
                                font_colors.append(str(run.font.color.rgb))
                        except (AttributeError, TypeError, ValueError):
                            pass
            shapes.append({
                "name": str(getattr(shape, "name", "")),
                "type": str(getattr(shape, "shape_type", "")),
                "left": int(shape.left),
                "top": int(shape.top),
                "width": int(shape.width),
                "height": int(shape.height),
                "text": str(getattr(shape, "text", "")),
                "editable_text": bool(getattr(shape, "has_text_frame", False)),
                "fill_color": fill_color,
                "font_colors": font_colors,
            })
        slides.append({"shape_count": len(shapes), "shapes": shapes})
    return {
        "slide_count": len(slides),
        "slide_width": int(presentation.slide_width),
        "slide_height": int(presentation.slide_height),
        "slides": slides,
    }


def compare_pptx_fixture(path: str | Path, fixture: Mapping[str, Any]) -> list[str]:
    """Compare an artifact to a small checked-in visual contract fixture."""

    snapshot = pptx_visual_snapshot(path)
    issues: list[str] = []
    expected_count = fixture.get("slide_count")
    if expected_count is not None and snapshot["slide_count"] != expected_count:
        issues.append(f"slide count changed: expected {expected_count}, got {snapshot['slide_count']}")
    expected_signature = fixture.get("signature")
    if expected_signature and pptx_visual_signature(path) != expected_signature:
        issues.append("PPT visual-contract signature changed")
    required_names = set(str(item) for item in fixture.get("required_shape_names", []))
    actual_names = {
        shape["name"]
        for slide in snapshot["slides"]
        for shape in slide["shapes"]
    }
    for name in sorted(required_names - actual_names):
        issues.append(f"required native layer is missing: {name}")
    min_editable = int(fixture.get("min_editable_text_shapes", 0))
    actual_editable = sum(
        int(shape["editable_text"])
        for slide in snapshot["slides"]
        for shape in slide["shapes"]
    )
    if actual_editable < min_editable:
        issues.append(f"editable text layer count fell below {min_editable}: {actual_editable}")
    return issues


def pptx_visual_signature(path: str | Path) -> str:
    """Hash geometry, ordering, and native-layer shape types, excluding copy."""

    snapshot = pptx_visual_snapshot(path)
    structural = {
        "slide_width": snapshot["slide_width"],
        "slide_height": snapshot["slide_height"],
        "slides": [
            {
                "shapes": [
                    {key: value for key, value in shape.items() if key != "text"}
                    for shape in slide["shapes"]
                ]
            }
            for slide in snapshot["slides"]
        ],
    }
    payload = json.dumps(structural, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


__all__ = ["compare_pptx_fixture", "pptx_visual_signature", "pptx_visual_snapshot"]
