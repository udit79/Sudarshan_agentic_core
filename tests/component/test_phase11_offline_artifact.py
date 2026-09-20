"""Proof that a real local source can reach a verified backend PPT artifact."""

from __future__ import annotations

from pptx import Presentation

from scripts.run_phase11_offline_artifact import build_offline_artifact


def test_offline_phase11_artifact_preserves_request_scope_and_provenance(tmp_path) -> None:
    source = tmp_path / "G01-normal-report.txt"
    source.write_text(
        "G01 Normal Report\n\n"
        "Verified observation 1: Three scheduled checks completed successfully.\n"
        "Unknown: The next review date is not provided.\n",
        encoding="utf-8",
    )

    result = build_offline_artifact(
        source,
        "Summarize the verified findings and list the unknowns.",
        output_root=tmp_path / "artifacts",
        user_id="fixture-operator",
        case_id="golden-g01",
        task_id="task-g01-offline-preview",
        source_reference="G01-normal-report.txt",
        run_id="run-g01-offline-preview",
    )

    assert result["status"] == "succeeded"
    assert result["provider_called"] is False
    assert result["memory_backend_called"] is False
    assert result["ingestion"]["indexed"] is True
    assert result["quality"] == {"approved": True, "slide_count": 4, "issues": []}

    manifest = result["artifact"]
    assert manifest["quality_status"] == "passed"
    assert manifest["user_id"] == "fixture-operator"
    assert manifest["case_id"] == "golden-g01"
    assert manifest["task_id"] == "task-g01-offline-preview"
    assert manifest["evidence_ids"]
    assert manifest["preview_uri"] is None

    artifact_path = tmp_path / "artifacts" / "artifacts" / "presentations"
    files = list(artifact_path.glob("*.pptx"))
    assert len(files) == 1
    presentation = Presentation(str(files[0]))
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert "Summarize the verified findings and list the unknowns." in text
    assert "Three scheduled checks completed successfully." in text
