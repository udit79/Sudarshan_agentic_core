from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation

from pipelines.ppt.ppt_master_adapter import PptMasterAdapter, PptMasterConfig


def test_ppt_master_export_reports_readback_slide_count(tmp_path: Path) -> None:
    root = tmp_path / "ppt-master"
    script = root / "skills" / "ppt-master" / "scripts" / "svg_to_pptx.py"
    script.parent.mkdir(parents=True)
    project = tmp_path / "project"
    (project / "svg_output").mkdir(parents=True)
    script.write_text(
        "from pathlib import Path\n"
        "import sys\n"
        "from pptx import Presentation\n"
        "project = Path(sys.argv[1])\n"
        "output = Path(sys.argv[sys.argv.index('-o') + 1])\n"
        "output.parent.mkdir(parents=True, exist_ok=True)\n"
        "prs = Presentation()\n"
        "prs.slides.add_slide(prs.slide_layouts[6])\n"
        "prs.slides.add_slide(prs.slide_layouts[6])\n"
        "prs.save(output)\n"
        "validation = project / 'validation'\n"
        "validation.mkdir(exist_ok=True)\n"
        "(validation / 'svg_quality_report.json').write_text('{}')\n"
        "(validation / f'{output.stem}.report.json').write_text('{}')\n",
        encoding="utf-8",
    )
    adapter = PptMasterAdapter(PptMasterConfig(root=root, python_executable=sys.executable, timeout_seconds=10))

    result = adapter.export(project, project / "exports" / "deck.pptx", roundtrip=True)

    assert result.status == "succeeded"
    assert result.slide_count == 2
    assert len(Presentation(result.output_path).slides) == 2
