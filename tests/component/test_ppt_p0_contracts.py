from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation

from pipelines.orchestrator.contracts import RequestConstraints
from pipelines.orchestrator.graph import _request_from_dict, _request_to_dict
from pipelines.common.contracts import AdvisoryRequest
from pipelines.ppt.presentation_quality import inspect_presentation
from pipelines.ppt.renderer import render_presentation
from pipelines.ppt.schemas import PresentationOutput, PresentationTheme, SlideContent, resolve_presentation_theme
from pipelines.ppt.visual_regression import compare_pptx_fixture


def _output() -> PresentationOutput:
    return PresentationOutput(
        presentation_id="p0-contract",
        title="P0 contract",
        classification_level="RESTRICTED",
        distribution="Authorized",
        slides=[
            SlideContent(slide_id="s1", order=1, title="Context", bullets=["Evidence"], layout="cover"),
            SlideContent(slide_id="s2", order=2, title="Decision", bullets=["Action"], layout="closing"),
        ],
    )


def test_page_count_is_canonical_and_conflicts_are_rejected() -> None:
    assert RequestConstraints(page_count=2).slide_count == 2
    with pytest.raises(ValueError, match="must match"):
        RequestConstraints(slide_count=2, page_count=3)


def test_constraints_survive_orchestrator_request_round_trip() -> None:
    request = AdvisoryRequest(
        query="make a two-page brief",
        user_id="u1",
        case_id="c1",
        task_id="t1",
        constraints=RequestConstraints(page_count=2, color_palette=["#112233"]),
    )
    restored = _request_from_dict(_request_to_dict(request))
    assert RequestConstraints.model_validate(restored.constraints).slide_count == 2


def test_exact_page_language_becomes_a_typed_constraint() -> None:
    request = AdvisoryRequest(
        query="Create exactly 2 pages about the case",
        user_id="u1",
        case_id="c1",
        task_id="t1",
    )
    assert RequestConstraints.model_validate(request.constraints).slide_count == 2
    only_request = AdvisoryRequest(
        query="I want only 2 pages in the PPT",
        user_id="u1",
        case_id="c1",
        task_id="t1",
    )
    assert RequestConstraints.model_validate(only_request.constraints).slide_count == 2

    plain_request = AdvisoryRequest(
        query="Make a PPT about the case with 2 slides",
        user_id="u1",
        case_id="c1",
        task_id="t1",
        requested_pipelines=("presentation",),
    )
    assert RequestConstraints.model_validate(plain_request.constraints).slide_count == 2

    source_pages = AdvisoryRequest(
        query="Summarize the first 2 pages of the source document",
        user_id="u1",
        case_id="c1",
        task_id="t1",
    )
    assert source_pages.constraints is None


def test_custom_theme_survives_native_ppt_render_and_fixture_gate(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    constraints = RequestConstraints(
        page_count=2,
        color_palette=["#112233", "#445566", "#778899"],
    )
    theme = resolve_presentation_theme(constraints)
    artifact = render_presentation(_output(), theme=theme)
    report = inspect_presentation(
        _output(),
        rendered_artifacts=[artifact.path],
        constraints=constraints,
    )
    assert report.approved, report.issues

    fixture = json.loads(
        (Path(__file__).parents[1] / "fixtures" / "ppt" / "ntro-briefing-v1.json").read_text(encoding="utf-8")
    )
    assert compare_pptx_fixture(artifact.path, fixture) == []

    presentation = Presentation(artifact.path)
    assert len(presentation.slides) == 2
    assert str(presentation.slides[0].background.fill.fore_color.rgb) == "0B1220"


def test_rendered_slide_count_is_checked_against_user_constraint(tmp_path, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    artifact = render_presentation(_output())
    report = inspect_presentation(
        _output(),
        rendered_artifacts=[artifact.path],
        constraints=RequestConstraints(page_count=3),
    )
    assert not report.approved
    assert any(issue.code == "PPT_CONSTRAINT_SLIDE_COUNT" for issue in report.issues)


def test_theme_contract_rejects_invalid_colors() -> None:
    with pytest.raises(ValueError, match="six-digit hex"):
        PresentationTheme(accent="red")
