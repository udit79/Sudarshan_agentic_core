from __future__ import annotations

from pptx import Presentation
import pytest

from pipelines.ppt.flowchart import (
    flowchart_from_visual_ir,
    layout_flowchart,
    render_flowchart_pptx,
    render_flowchart_svg,
)
from pipelines.ppt.schemas import FlowchartEdge, FlowchartNode, FlowchartSpec, LayoutBox, VisualIR


def spec() -> FlowchartSpec:
    return FlowchartSpec(
        flowchart_id="case-process",
        nodes=[
            FlowchartNode(node_id="ingest", label="Ingest", kind="start"),
            FlowchartNode(node_id="review", label="Review", kind="review"),
            FlowchartNode(node_id="decision", label="Decision", kind="decision"),
            FlowchartNode(node_id="release", label="Release", kind="end"),
        ],
        edges=[
            FlowchartEdge(source="ingest", target="review", label="source"),
            FlowchartEdge(source="review", target="decision"),
            FlowchartEdge(source="decision", target="release", label="approved"),
        ],
    )


def test_layout_and_svg_share_all_nodes_and_edges() -> None:
    chart = spec()
    layout = layout_flowchart(chart)
    svg = render_flowchart_svg(chart, layout=layout)

    assert len(layout.nodes) == 4
    assert all(0 <= item.box.x <= 1 and 0 <= item.box.y <= 1 for item in layout.nodes)
    assert svg.startswith("<svg")
    assert svg.count("<rect") == 5  # background plus four editable node previews
    assert "Ingest" in svg and "approved" in svg
    assert "marker-end" in svg


def test_flowchart_rejects_cycles_and_isolated_nodes() -> None:
    cyclic = FlowchartSpec(
        flowchart_id="cycle",
        nodes=[FlowchartNode(node_id="a", label="A"), FlowchartNode(node_id="b", label="B")],
        edges=[FlowchartEdge(source="a", target="b"), FlowchartEdge(source="b", target="a")],
    )
    with pytest.raises(ValueError, match="acyclic"):
        layout_flowchart(cyclic)

    isolated = FlowchartSpec(
        flowchart_id="isolated",
        nodes=[FlowchartNode(node_id="a", label="A"), FlowchartNode(node_id="b", label="B")],
    )
    with pytest.raises(ValueError, match="isolated"):
        layout_flowchart(isolated)


def test_pptx_renderer_emits_editable_shapes_and_connectors(tmp_path) -> None:
    path = tmp_path / "flowchart.pptx"
    artifact = render_flowchart_pptx(spec(), path, title="Case process")
    presentation = Presentation(str(path))
    slide = presentation.slides[0]

    assert artifact.node_count == 4
    assert artifact.edge_count == 3
    assert len(presentation.slides) == 1
    assert sum(1 for shape in slide.shapes if shape.shape_type == 1) >= 4  # AUTO_SHAPE
    assert sum(1 for shape in slide.shapes if shape.shape_type == 9) == 3  # CONNECTOR
    assert "Release" in " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))


def test_visual_ir_adapter_preserves_typed_graph_boundary() -> None:
    visual = VisualIR(
        visual_id="visual-1",
        kind="flowchart",
        bounds=LayoutBox(x=0.1, y=0.1, width=0.8, height=0.8),
        alt_text="A simple process",
        data={
            "nodes": [{"id": "a", "label": "A"}, {"id": "b", "label": "B", "kind": "end"}],
            "edges": [{"from": "a", "to": "b"}],
        },
    )
    converted = flowchart_from_visual_ir(visual)
    assert converted.flowchart_id == "visual-1"
    assert converted.edges[0].source == "a"
